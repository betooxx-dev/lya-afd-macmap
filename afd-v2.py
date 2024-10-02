import tkinter as tk
from tkinter import ttk, filedialog
import csv
import os
import openpyxl
from docx import Document
from bs4 import BeautifulSoup

transitions = {
    ('q0', 'hex'): 'q1',
    ('q0', 'space'): 'q0',
    ('q1', 'hex'): 'q2',
    ('q2', ':'): 'q3',
    ('q2', '-'): 'q4',
    ('q2', 'space'): 'q5',
    ('q3', 'hex'): 'q6',
    ('q4', 'hex'): 'q7',
    ('q5', 'hex'): 'q8',
    ('q6', 'hex'): 'q9',
    ('q7', 'hex'): 'q10',
    ('q8', 'hex'): 'q11',
    ('q9', ':'): 'q12',
    ('q10', '-'): 'q13',
    ('q11', 'space'): 'q14',
    ('q12', 'hex'): 'q15',
    ('q13', 'hex'): 'q16',
    ('q14', 'hex'): 'q17',
    ('q15', 'hex'): 'q18',
    ('q16', 'hex'): 'q19',
    ('q17', 'hex'): 'q20',
    ('q18', ':'): 'q21',
    ('q19', '-'): 'q22',
    ('q20', 'space'): 'q23',
    ('q21', 'hex'): 'q24',
    ('q22', 'hex'): 'q25',
    ('q23', 'hex'): 'q26',
    ('q24', 'hex'): 'q27',
    ('q25', 'hex'): 'q28',
    ('q26', 'hex'): 'q29',
    ('q27', ':'): 'q30',
    ('q28', '-'): 'q31',
    ('q29', 'space'): 'q32',
    ('q30', 'hex'): 'q33',
    ('q31', 'hex'): 'q34',
    ('q32', 'hex'): 'q35',
    ('q33', 'hex'): 'q36',
    ('q34', 'hex'): 'q37',
    ('q35', 'hex'): 'q38',
    ('q36', ':'): 'q39',
    ('q37', '-'): 'q40',
    ('q38', 'space'): 'q41',
    ('q39', 'hex'): 'q42',
    ('q40', 'hex'): 'q42',
    ('q41', 'hex'): 'q42',
    ('q42', 'hex'): 'q43',
    ('q43', 'space'): 'q44',
    ('q43', 'end'): 'q44'
}

class MacAFD:
    def __init__(self):
        self.state = 'q0'
        self.transitions = transitions

    def is_hex(self, char):
        return char.lower() in '0123456789abcdef'

    def transition(self, char):
        if char.isspace():
            input_type = 'space'
        elif self.is_hex(char):
            input_type = 'hex'
        elif char in [':', '-']:
            input_type = char
        elif not char:
            input_type = 'end'
        else:
            input_type = 'other'

        next_state = self.transitions.get((self.state, input_type), 'q0')
        self.state = next_state

    def is_accepted(self):
        return self.state == 'q44'

def find_valid_macs(text):
    afd = MacAFD()
    valid_macs = []
    current_mac = ""
    start_index = 0

    for i, char in enumerate(text):
        prev_state = afd.state
        afd.transition(char)
        
        if prev_state == 'q0' and afd.state == 'q1':
            start_index = i
            current_mac = char
        elif afd.state != 'q0':
            current_mac += char
        
        if afd.is_accepted():
            valid_macs.append((current_mac.strip(), start_index))
            current_mac = ""
            afd = MacAFD()  # Reiniciamos el afd

    # Manejar el caso del final de la cadena
    afd.transition('')
    if afd.is_accepted():
        valid_macs.append((current_mac.strip(), start_index))

    return valid_macs

class MacRecognizer:
    def __init__(self):
        self.window = tk.Tk()
        self.window.title("MacMap - Reconocedor de direcciones MAC")
        self.window.geometry("650x550")
        self.window.configure(bg="#f0f0f0")

        self.style = ttk.Style()
        self.style.theme_use('clam')
        self.configure_styles()

        self.create_widgets()

    def configure_styles(self):
        self.style.configure('TFrame', background="#f0f0f0")
        self.style.configure('TLabel', background="#f0f0f0", font=('Arial', 12))
        self.style.configure('TButton', font=('Arial', 12, 'bold'), background="#4CAF50", foreground="white")
        self.style.map('TButton', background=[('active', '#45a049')])

    def create_widgets(self):
        main_frame = ttk.Frame(self.window, padding="20")
        main_frame.grid(row=0, column=0, sticky=(tk.W, tk.E, tk.N, tk.S))
        self.window.grid_columnconfigure(0, weight=1)
        self.window.grid_rowconfigure(0, weight=1)

        title_label = ttk.Label(main_frame, text="Reconocedor de direcciones MAC", font=('Arial', 16, 'bold'))
        title_label.grid(row=0, column=0, columnspan=2, pady=(0, 20))

        ttk.Label(main_frame, text="Seleccione el archivo de entrada:").grid(row=1, column=0, sticky=tk.W, pady=5)
        self.file_path = tk.StringVar()
        self.file_entry = ttk.Entry(main_frame, textvariable=self.file_path, width=50)
        self.file_entry.grid(row=2, column=0, pady=5, padx=(0, 10))
        self.browse_button = ttk.Button(main_frame, text="Examinar", command=self.browse_file)
        self.browse_button.grid(row=2, column=1, pady=5)

        self.start_button = ttk.Button(main_frame, text="Empezar Análisis", command=self.start_analysis)
        self.start_button.grid(row=3, column=0, columnspan=2, pady=15)

        self.status_label = ttk.Label(main_frame, text="", font=('Arial', 11, 'italic'))
        self.status_label.grid(row=4, column=0, columnspan=2, pady=5)

        result_frame = ttk.Frame(main_frame, padding="10")
        result_frame.grid(row=5, column=0, columnspan=2, sticky=(tk.W, tk.E, tk.N, tk.S), pady=10)
        result_frame.grid_columnconfigure(0, weight=1)
        result_frame.grid_rowconfigure(0, weight=1)

        self.result_text = tk.Text(result_frame, height=12, width=70, font=('Arial', 11), bg="white", state='disabled')
        self.result_text.grid(row=0, column=0, sticky=(tk.W, tk.E, tk.N, tk.S))

        scrollbar = ttk.Scrollbar(result_frame, orient="vertical", command=self.result_text.yview)
        scrollbar.grid(row=0, column=1, sticky=(tk.N, tk.S))
        self.result_text.configure(yscrollcommand=scrollbar.set)

    def browse_file(self):
        filetypes = [
            ('Archivos soportados', '*.xlsx *.csv *.docx *.html *.txt'),
            ('Excel', '*.xlsx'),
            ('CSV', '*.csv'),
            ('Word', '*.docx'),
            ('HTML', '*.html'),
            ('TXT', '*.txt'),
            ('Todos los archivos', '*.*')
        ]
        filename = filedialog.askopenfilename(filetypes=filetypes)
        if filename:
            self.file_path.set(filename)

    def read_file(self, file_path):
        _, file_extension = os.path.splitext(file_path)
        
        if file_extension == '.xlsx':
            return self.read_excel(file_path)
        elif file_extension == '.csv':
            return self.read_csv(file_path)
        elif file_extension == '.docx':
            return self.read_docx(file_path)
        elif file_extension == '.html':
            return self.read_html(file_path)
        elif file_extension == '.txt':
            return self.read_txt(file_path)
        else:
            raise ValueError("Formato de archivo no soportado")

    def read_excel(self, file_path):
        workbook = openpyxl.load_workbook(file_path)
        sheet = workbook.active
        content = []
        for row in sheet.iter_rows(values_only=True):
            content.append([str(cell) if cell is not None else '' for cell in row])
        return content

    def read_csv(self, file_path):
        with open(file_path, 'r', newline='', encoding='utf-8') as csvfile:
            reader = csv.reader(csvfile)
            return list(reader)

    def read_docx(self, file_path):
        doc = Document(file_path)
        return [[paragraph.text] for paragraph in doc.paragraphs]

    def read_html(self, file_path):
        with open(file_path, 'r', encoding='utf-8') as htmlfile:
            soup = BeautifulSoup(htmlfile, 'html.parser')
            return [[element.get_text()] for element in soup.find_all(['p', 'div', 'span', 'h1', 'h2', 'h3', 'h4', 'h5', 'h6'])]

    def read_txt(self, file_path):
        with open(file_path, 'r', encoding='utf-8') as txtfile:
            return [line for line in txtfile]

    def start_analysis(self):
        self.start_button.config(state='disabled')
        self.status_label.config(text="Analizando...", foreground="blue")
        self.window.update()

        file_path = self.file_path.get()
        if not file_path:
            self.status_label.config(text="Por favor, seleccione un archivo", foreground="red")
            self.start_button.config(state='normal')
            return

        try:
            content = self.read_file(file_path)
            valid_macs = self.find_macs_in_content(content)
            self.show_results(valid_macs)
            self.save_results_to_csv(valid_macs)
            self.status_label.config(text="Análisis completado. Resultados guardados en 'resultados_mac.csv'", foreground="green")
        except Exception as e:
            self.status_label.config(text=f"Error: {str(e)}", foreground="red")
        finally:
            self.start_button.config(state='normal')

    def find_macs_in_content(self, content):
        valid_macs = []
        for row_num, row in enumerate(content, start=1):
            if isinstance(row, list):
                for col_num, cell in enumerate(row, start=1):
                    macs = find_valid_macs(cell)
                    for mac, position in macs:
                        valid_macs.append((mac, row_num, col_num, position))
            else:
                macs = find_valid_macs(row)
                for mac, position in macs:
                    valid_macs.append((mac, row_num, 1, position))
        return valid_macs

    def show_results(self, valid_macs):
        self.result_text.config(state='normal')
        self.result_text.delete("1.0", tk.END)
        if valid_macs:
            for i, (mac, row, col, position) in enumerate(valid_macs, 1):
                self.result_text.insert(tk.END, f"{i}. MAC válida: ", "bold")
                self.result_text.insert(tk.END, f"{mac} ", "mac")
                self.result_text.insert(tk.END, f"(Fila: {row}, Columna: {col}, Posición: {position})\n")
        else:
            self.result_text.insert(tk.END, "No se encontraron direcciones MAC válidas.", "italic")
        self.result_text.config(state='disabled')

        self.result_text.tag_configure("bold", font=('Arial', 11, 'bold'))
        self.result_text.tag_configure("mac", foreground="blue")
        self.result_text.tag_configure("italic", font=('Arial', 11, 'italic'))

    def save_results_to_csv(self, valid_macs):
        with open('resultados_mac.csv', 'w', newline='', encoding='utf-8') as csvfile:
            writer = csv.writer(csvfile)
            writer.writerow(['MAC', 'Fila', 'Columna', 'Posición'])
            for mac, row, col, position in valid_macs:
                writer.writerow([mac, row, col, position])

    def run(self):
        self.window.mainloop()

if __name__ == "__main__":
    app = MacRecognizer()
    app.run()